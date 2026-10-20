"""Build the six invitation-request drafts (DOCX + PDF) and the combined PDF.

No organizer of these three events publishes an invitation-letter template, so
every document is an independently prepared *request* written in the
attendee's voice. Nothing here is issued, signed or confirmed by an organizer.

Requires: python-docx, pypdf, LibreOffice (soffice) on PATH, and the fonts
Inter, Roboto, DM Sans and Bebas Neue installed for faithful PDF output.

Usage:  python3 build_documents.py            (writes into ../docx and ../pdf)
"""

import os
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader, PdfWriter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_DIR = os.path.join(ROOT, "docx")
PDF_DIR = os.path.join(ROOT, "pdf")
CHECKED = "20 October 2026"
BLANK = "______________________________"

ATTENDEES = [
    {"name": "Sebastian-Eugen Buzdugan", "slug": "Sebastian-Eugen-Buzdugan", "first": "Sebastian-Eugen"},
    {"name": "Codruta Costea", "slug": "Codruta-Costea", "first": "Codruta"},
]

EVENTS = [
    {
        "key": "01_AI-Engineers-Meetup-Bangkok",
        "title": "AI Engineers Meetup Bangkok",
        "short": "AI Engineers Meetup",
        "kicker": "AI ENGINEERS IN THAILAND · MEETUP",
        "when": "Saturday, 26 December 2026 · 10:00–13:00 ICT",
        "where": "SCBX NEXT TECH, 4th Floor, Siam Paragon, Bangkok, Thailand",
        "attend": "Saturday, 26 December 2026",
        "organizer": "AI Engineers in Thailand (Meetup group), host Ek C.",
        "organizer_short": "the AI Engineers in Thailand Meetup group",
        "doc_title": "Request for Attendance Confirmation",
        "to_lines": [
            "Ek C., Event Host",
            "AI Engineers in Thailand (Meetup group “coding-in-bangkok”)",
            "Via Meetup group messaging",
        ],
        "salutation": "Dear Ek,",
        "subject": "Attendance confirmation — AI Engineers Meetup Bangkok, 26 December 2026",
        "body": [
            "I would like to attend the AI Engineers Meetup Bangkok on Saturday, 26 December 2026 "
            "(10:00–13:00 ICT) at SCBX NEXT TECH, 4th Floor, Siam Paragon. I will be in Bangkok in late "
            "December 2026, and the weekly session's focus on trying out the latest AI tools is closely "
            "related to my professional interests.",
            "Because I am traveling internationally, I would be grateful for a short letter, or a written "
            "message through Meetup or by email, confirming the event details below and that I am welcome "
            "to attend. I understand this is a free community meetup and that the group may not normally "
            "provide letters, so a brief written confirmation of these facts would be entirely sufficient.",
            "I will cover all of my own travel, accommodation and other costs. I am asking only for a "
            "statement of the facts listed below, not for any sponsorship or financial undertaking.",
        ],
        "status_label": "Meetup RSVP status",
        "status_options": "☐ RSVP submitted on Meetup     ☐ RSVP not yet submitted",
        "fields_extra": [],
        "facts": [
            ("Event", "AI Engineers Meetup Bangkok (weekly series; 26 Dec 2026 is the last listed date)"),
            ("Date and time", "Saturday, 26 December 2026, 10:00–13:00 ICT"),
            ("Venue", "SCBX NEXT TECH, 4th Fl., Siam Paragon, Bangkok (no street address on the listing)"),
            ("Organizer", "Meetup group “AI Engineers in Thailand”; host shown as Ek C."),
            ("Fee", "None stated on the listing"),
            ("Registration", "RSVP through Meetup; no attendee count, capacity or invitation-letter process published"),
        ],
        "sources": [
            ("Official event listing", "https://www.meetup.com/th-th/coding-in-bangkok/events/nsntvtyjcqbjc/"),
            ("Organizer group page", "https://www.meetup.com/coding-in-bangkok/"),
        ],
        "outstanding": [
            "RSVP to the 26 December 2026 session from your own Meetup account, and check that the session is still scheduled shortly before traveling (it is the final date of a recurring series).",
            "Send this request to the host through Meetup messaging; the group publishes no invitation-letter procedure or contact email.",
            "Check Thailand's entry requirements for your nationality, including the Thailand Digital Arrival Card, with official Thai government sources.",
        ],
        "brand_note": "The organizer publishes no brand assets beyond its Meetup listing, so this draft uses a neutral design.",
        "style": {
            "body_font": "Inter", "display_font": "Inter", "display_bold": True, "display_size": 21,
            "band": "1F2937", "band_text": "FFFFFF", "band_sub": "CBD5E1",
            "accent": "0F766E", "heading": "111827", "text": "1F2937", "muted": "6B7280",
            "link": "0F766E", "rule": "0F766E", "label": "0F766E", "fill": "F3F4F6",
        },
    },
    {
        "key": "02_AI-for-Science-Congress-2026",
        "title": "AI for Science Congress 2026",
        "short": "AI for Science Congress 2026",
        "kicker": "ISPA · BDCLOUD · SOCIALCOM · SUSTAINCOM · SPACCS 2026",
        "when": "27–30 December 2026 (attending 29–30 December)",
        "where": "Kuala Lumpur, Malaysia (venue not yet published)",
        "attend": "Tuesday 29 – Wednesday 30 December 2026",
        "organizer": "AI for Science Congress 2026 Organizing Committee",
        "organizer_short": "the AI for Science Congress 2026 organizers",
        "doc_title": "Request for Invitation Letter",
        "to_lines": [
            "Organizing Committee, AI for Science Congress 2026",
            "Kuala Lumpur, Malaysia",
            "ai4science-2026@googlegroups.com",
        ],
        "salutation": "Dear Members of the Organizing Committee,",
        "subject": "Attendance and invitation letter request — AI for Science Congress 2026, 29–30 December 2026",
        "body": [
            "I am writing to ask about attending the AI for Science Congress 2026 in Kuala Lumpur as a "
            "non-presenting attendee on 29 and 30 December 2026, and to request a formal invitation letter "
            "once my registration is complete. The congress's focus on artificial intelligence and its "
            "co-located conferences is closely related to my professional interests.",
            "The registration instructions currently published on the congress website describe registration "
            "for paper authors. Could you please confirm whether attendee (non-author) registration is "
            "available for these dates, the applicable fee and payment procedure, and how invitation letters "
            "are requested?",
            "Once I am registered, I would be grateful for a letter stating my name as it appears in my "
            "passport, the congress title, dates and location, my dates of attendance, and my registration "
            "status. I will cover all of my own travel, accommodation and registration costs.",
        ],
        "status_label": "Registration status",
        "status_options": "☐ Registered and paid     ☐ Awaiting organizer reply on attendee registration",
        "fields_extra": [("Registration ID (issued by organizer)", BLANK)],
        "facts": [
            ("Event", "AI for Science Congress 2026, with ISPA, BDCloud, SocialCom, SustainCom and SpaCCS 2026"),
            ("Dates", "27–30 December 2026 (requested attendance: 29–30 December)"),
            ("Location", "Kuala Lumpur, Malaysia; the venue page still shows 2025 (Zhengzhou) content"),
            ("Organizers named", "“Sponsored and Organized by” Zhengzhou University and Zhejiang Normal University"),
            ("Registration", "Published only for paper authors (Regular registration, deadline 30 Nov 2026)"),
            ("Contact", "ai4science-2026@googlegroups.com"),
        ],
        "sources": [
            ("Official website", "https://ieee-ai-for-science.org/2026/"),
            ("Contact page", "https://ieee-ai-for-science.org/2026/contact.php"),
            ("Registration instructions", "https://ieee-ai-for-science.org/2026/ispa/author/aure.php"),
            ("Venue page", "https://ieee-ai-for-science.org/2026/venue.php"),
        ],
        "outstanding": [
            "Ask the organizers whether non-author attendee registration exists for 29–30 December 2026. The published fees (for example USD 750 IEEE member / USD 900 non-member Regular registration) are for paper authors.",
            "The congress publishes no invitation-letter procedure; ask for one at ai4science-2026@googlegroups.com and have the letter issued only on the organizers' own letterhead.",
            "The venue, program and travel pages are not yet published for 2026 (the travel page returns an error). Confirm the venue address before booking.",
            "Payment is by bank transfer per the registration page; confirm bank details directly with the organizers by email before paying.",
            "Check Malaysia's entry requirements for your nationality and complete the Malaysia Digital Arrival Card (MDAC) before arrival.",
        ],
        "brand_note": "Styled after the congress website (black header, Roboto). No organizer or IEEE logos are used, because this is not an organizer document.",
        "style": {
            "body_font": "Roboto", "display_font": "Roboto", "display_bold": True, "display_size": 21,
            "band": "000000", "band_text": "FFFFFF", "band_sub": "BFBFBF",
            "accent": "0D6EFD", "heading": "141414", "text": "333333", "muted": "6C757D",
            "link": "0D6EFD", "rule": "141414", "label": "0D6EFD", "fill": "F2F2F2",
        },
    },
    {
        "key": "03_MAIXPO-2026",
        "title": "MAIXPO 2026",
        "short": "MAIXPO 2026",
        "kicker": "AI MARKETING CONFERENCE & EXPO",
        "when": "29–30 December 2026",
        "where": "Kuala Lumpur, Malaysia (venue to be announced)",
        "attend": "Tuesday 29 – Wednesday 30 December 2026",
        "organizer": "MAIXPO (organizing entity not named on the website)",
        "organizer_short": "MAIXPO",
        "doc_title": "Request for Visa Invitation Letter",
        "to_lines": [
            "MAIXPO Team",
            "MAIXPO 2026 — AI Marketing Conference & Expo, Kuala Lumpur",
            "info@maixpo.com",
        ],
        "salutation": "Dear MAIXPO Team,",
        "subject": "Visa invitation letter request — MAIXPO 2026, Kuala Lumpur, 29–30 December 2026",
        "body": [
            "Following the procedure published on the MAIXPO website, I am requesting a formal invitation "
            "letter for MAIXPO 2026 — AI Marketing Conference & Expo, taking place in Kuala Lumpur, Malaysia, "
            "on 29–30 December 2026. I intend to attend in person on both days.",
            "The details your website asks for (full name, passport number, nationality and country of "
            "application) are set out below, together with my ticket details. I understand that invitation "
            "letters are issued only to confirmed ticket holders.",
            "As the venue has not yet been announced, I would also be grateful if the letter, or a follow-up "
            "message, could include the venue address once it is confirmed. I will cover all of my own "
            "travel, accommodation and ticket costs.",
        ],
        "status_label": "Ticket status",
        "status_options": "☐ Ticket confirmed and paid     ☐ Registered, payment pending",
        "fields_extra": [
            ("Ticket type", "☐ Standard     ☐ VIP / Executive"),
            ("Registration reference (from MAIXPO)", BLANK),
        ],
        "facts": [
            ("Event", "MAIXPO 2026 — AI Marketing Conference & Expo (hybrid: in person and virtual)"),
            ("Dates", "29–30 December 2026; the site calls these “Wednesday & Thursday”, but they fall on Tuesday and Wednesday"),
            ("Location", "Kuala Lumpur, Malaysia; venue “will be announced once confirmed”"),
            ("Tickets", "Standard early bird USD 49, VIP early bird USD 99, both until 30 Oct 2026; online payment “coming soon”"),
            ("Letter policy", "Issued to confirmed ticket holders; email full name, passport number, nationality and country of application"),
            ("Contact", "info@maixpo.com (general inquiries and invitation letters)"),
        ],
        "sources": [
            ("Official website", "https://maixpo.com/"),
            ("Tickets and registration", "https://maixpo.com/#tickets"),
            ("Invitation-letter policy (FAQ)", "https://maixpo.com/#faq"),
        ],
        "outstanding": [
            "Register and obtain a confirmed ticket first; letters go only to confirmed ticket holders. Early-bird pricing ends 30 October 2026, and payment instructions arrive by email because online payment is not yet live.",
            "Email this request to info@maixpo.com with your passport number, nationality and country of application filled in.",
            "Confirm the weekday mismatch (the site says Wednesday & Thursday; 29–30 December 2026 are Tuesday and Wednesday) and the venue address, which is still to be announced.",
            "Before paying, note that the website names no legal entity or company address, its Privacy and Terms links open an email instead of a policy, and its speakers are listed as TBA.",
            "Check Malaysia's entry requirements for your nationality and complete the Malaysia Digital Arrival Card (MDAC) before arrival.",
        ],
        "brand_note": "Styled after the MAIXPO website palette (#080808, #E8FF00, #FF4D00) and typefaces (Bebas Neue, DM Sans). The MAIXPO logo is not used, because this is not an organizer document.",
        "style": {
            "body_font": "DM Sans", "display_font": "Bebas Neue", "display_bold": False, "display_size": 34,
            "band": "080808", "band_text": "F5F2EC", "band_sub": "E8FF00",
            "accent": "FF4D00", "heading": "080808", "text": "1A1A1A", "muted": "5C5C5C",
            "link": "C23A00", "rule": "FF4D00", "label": "FF4D00", "fill": "F5F2EC",
        },
    },
]


# ---------- low-level docx helpers ----------

def rgb(hex6):
    return RGBColor.from_string(hex6)


def shade(cell, hex6):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    tc_pr.append(shd)


def cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(val * 56.7)))  # mm -> twips
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def table_borders(table, color=None, size=4, inside=True):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ())
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if color and edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def bottom_border(cell, color, size=8):
    tc_pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:color"), color)
    b.append(el)
    tc_pr.append(b)


def fixed_widths(table, widths_mm):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_mm):
        col.set(qn("w:w"), str(int(w * 56.7)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)


def para(container, text="", font=None, size=10, bold=False, italic=False, color="000000",
         align=None, before=0, after=4, line=1.2, keep_next=False):
    p = container.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    if align:
        p.alignment = align
    if text:
        run(p, text, font, size, bold, italic, color)
    return p


def run(p, text, font=None, size=10, bold=False, italic=False, color="000000"):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(color)
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        r._element.rPr.rFonts.set(qn("w:cs"), font)
    return r


def hyperlink(p, url, text, font, size, color):
    rel = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), font)
    r_pr.append(fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def field(p, instr, font, size, color):
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        r = run(p, "", font, size, color=color)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._element.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            r._element.append(it)
        else:
            r.text = text


# ---------- document builder ----------

def build(event, attendee, companion):
    s = event["style"]
    bf = s["body_font"]
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = bf
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), bf)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(16), Mm(18)
    sec.header_distance, sec.footer_distance = Mm(8), Mm(9)
    content_w = 170

    # footer: status line + page numbers on every page
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)
    run(fp, f"Draft request by {attendee['name']} · Not issued, signed or confirmed by the organizer · Page ",
        bf, 7.5, color=s["muted"])
    field(fp, "PAGE", bf, 7.5, s["muted"])
    run(fp, " of ", bf, 7.5, color=s["muted"])
    field(fp, "NUMPAGES", bf, 7.5, s["muted"])

    # 1. draft status banner
    t = doc.add_table(rows=1, cols=1)
    table_borders(t, "B45309", 6)
    fixed_widths(t, [content_w])
    c = t.cell(0, 0)
    shade(c, "FEF3C7")
    cell_margins(c, 1.6, 1.6, 3, 3)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    run(p, "INDEPENDENTLY PREPARED DRAFT — NOT AN OFFICIAL ORGANIZER DOCUMENT", bf, 8.5, True, color="78350F")
    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run(p, f"Written by the attendee to request a letter from {event['organizer_short']}. It does not show that "
           "registration, payment, a ticket or an invitation has been confirmed or issued.", bf, 7.5,
        color="78350F")

    para(doc, "", size=4, after=0)

    # 2. event header band
    t = doc.add_table(rows=1, cols=1)
    table_borders(t)
    fixed_widths(t, [content_w])
    c = t.cell(0, 0)
    shade(c, s["band"])
    cell_margins(c, 4, 4, 6, 6)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    run(p, event["kicker"], bf, 7.5, True, color=s["band_sub"])
    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run(p, event["title"], s["display_font"], s["display_size"], s["display_bold"], color=s["band_text"])
    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run(p, event["when"] + "  ·  " + event["where"], bf, 8.5, color=s["band_text"])
    # accent rule under the band
    t2 = doc.add_table(rows=1, cols=1)
    table_borders(t2)
    fixed_widths(t2, [content_w])
    c2 = t2.cell(0, 0)
    shade(c2, s["rule"])
    c2.paragraphs[0].paragraph_format.space_after = Pt(0)
    run(c2.paragraphs[0], " ", bf, 2)
    c2.paragraphs[0].paragraph_format.line_spacing = Pt(3)

    # 3. title
    para(doc, event["doc_title"], s["display_font"] if s["display_font"] != "Bebas Neue" else bf,
         15, True, color=s["heading"], before=8, after=5)

    # 4. from / to block
    t = doc.add_table(rows=1, cols=2)
    table_borders(t)
    fixed_widths(t, [85, 85])
    left, right = t.cell(0, 0), t.cell(0, 1)
    for cell, label, lines in (
        (left, "FROM", [attendee["name"], "Address: " + BLANK[:22], "Email: " + BLANK[:24],
                        "Phone: " + BLANK[:24]]),
        (right, "TO", event["to_lines"]),
    ):
        cell_margins(cell, 0, 0, 0, 2)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        run(p, label, bf, 7, True, color=s["label"])
        for i, line in enumerate(lines):
            q = cell.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            run(q, line, bf, 9, i == 0, color=s["text"])
    p = para(doc, "", before=5, after=5)
    run(p, "Date: ", bf, 9, True, color=s["text"])
    run(p, BLANK[:20], bf, 9, color=s["text"])

    # 5. subject + body
    p = para(doc, "", after=6)
    run(p, "Subject: ", bf, 10, True, color=s["heading"])
    run(p, event["subject"], bf, 10, True, color=s["heading"])
    para(doc, event["salutation"], bf, 9.5, color=s["text"], after=5)
    for i, text in enumerate(event["body"]):
        if i == 1:
            text += (f" I will be traveling with {companion['name']}, who is submitting a separate "
                     "request under their own name.")
        para(doc, text, bf, 9.5, color=s["text"], after=5, line=1.18)

    # 6. details for the letter
    para(doc, "Details for the letter", bf, 10.5, True, color=s["heading"], before=4, after=3, keep_next=True)
    rows = [
        ("Full name (exactly as in passport)", attendee["name"]),
        ("Nationality", BLANK),
        ("Passport number", BLANK),
        ("Passport expiry date", BLANK),
        ("Country of visa application", BLANK),
        ("Organization and position", BLANK),
        ("Dates of attendance", event["attend"]),
        (event["status_label"], event["status_options"]),
    ] + event["fields_extra"]
    t = doc.add_table(rows=len(rows), cols=2)
    table_borders(t, "D0D0D0", 4)
    fixed_widths(t, [62, 108])
    for i, (k, v) in enumerate(rows):
        a, b = t.cell(i, 0), t.cell(i, 1)
        shade(a, s["fill"])
        for cell in (a, b):
            cell_margins(cell, 0.7, 0.7, 2, 2)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(a.paragraphs[0], k, bf, 8.5, True, color=s["text"])
        run(b.paragraphs[0], v, bf, 8.5, v == attendee["name"], color=s["text"])
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    # 7. closing + requester signature (the attendee signs their own request)
    para(doc, "Thank you for your time and assistance.  Yours sincerely,", bf, 9.5, color=s["text"],
         before=7, after=18, keep_next=True)
    p = para(doc, "_________________________________", bf, 10, color=s["muted"], after=1, keep_next=True)
    para(doc, attendee["name"] + " (requester)", bf, 10, True, color=s["text"], after=0)

    # ---------- page 2: annex ----------
    p = para(doc, "Annex — Event details and status", s["display_font"] if s["display_font"] != "Bebas Neue" else bf,
         14, True, color=s["heading"], after=2)
    p.paragraph_format.page_break_before = True
    para(doc, f"Facts below are taken only from the organizer's public pages, checked on {CHECKED}. "
              "They are provided so the organizer can verify the request quickly.",
         bf, 8.5, italic=True, color=s["muted"], after=6)

    t = doc.add_table(rows=len(event["facts"]), cols=2)
    table_borders(t, "D0D0D0", 4)
    fixed_widths(t, [38, 132])
    for i, (k, v) in enumerate(event["facts"]):
        a, b = t.cell(i, 0), t.cell(i, 1)
        shade(a, s["fill"])
        for cell in (a, b):
            cell_margins(cell, 1, 1, 2, 2)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        run(a.paragraphs[0], k, bf, 8.5, True, color=s["text"])
        run(b.paragraphs[0], v, bf, 8.5, color=s["text"])

    para(doc, "Official sources", bf, 10.5, True, color=s["heading"], before=10, after=3)
    for label, url in event["sources"]:
        p = para(doc, "", after=2)
        run(p, label + ": ", bf, 8.5, True, color=s["text"])
        hyperlink(p, url, url, bf, 8.5, s["link"])

    para(doc, "Outstanding before this request is sent", bf, 10.5, True, color=s["heading"], before=10, after=3)
    for item in event["outstanding"]:
        p = para(doc, "", after=3, line=1.2)
        p.paragraph_format.left_indent = Mm(5)
        p.paragraph_format.first_line_indent = Mm(-5)
        run(p, "☐ ", bf, 9, color=s["label"])
        run(p, item, bf, 9, color=s["text"])

    para(doc, "Document status", bf, 10.5, True, color=s["heading"], before=10, after=3)
    t = doc.add_table(rows=1, cols=1)
    table_borders(t, "B45309", 6)
    fixed_widths(t, [content_w])
    c = t.cell(0, 0)
    shade(c, "FFFBEB")
    cell_margins(c, 2, 2, 3, 3)
    lines = [
        ("Template: ", "no organizer-provided invitation-letter template was found for this event. This is an "
                       "independently designed request, not a reproduction of an official format."),
        ("Not included: ", "organizer signatures, stamps, seals, registration or ticket numbers, QR codes and "
                           "payment confirmations. Any official letter must come from the organizer."),
        ("Design: ", event["brand_note"]),
    ]
    for i, (k, v) in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(2 if i < len(lines) - 1 else 0)
        run(p, k, bf, 8.5, True, color="78350F")
        run(p, v, bf, 8.5, color="78350F")

    core = doc.core_properties
    core.title = f"{event['doc_title']} — {event['title']} — {attendee['name']}"
    core.author = attendee["name"]
    core.subject = "Independently prepared draft request; not an official organizer document"
    core.comments = f"Facts checked {CHECKED}."
    return doc


def main():
    os.makedirs(DOCX_DIR, exist_ok=True)
    os.makedirs(PDF_DIR, exist_ok=True)
    built = []
    for event in EVENTS:
        for idx, attendee in enumerate(ATTENDEES):
            companion = ATTENDEES[1 - idx]
            name = f"{event['key']}_{attendee['slug']}"
            path = os.path.join(DOCX_DIR, name + ".docx")
            build(event, attendee, companion).save(path)
            built.append((event, attendee, name))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", PDF_DIR]
                   + [os.path.join(DOCX_DIR, n + ".docx") for _, _, n in built],
                   check=True, stdout=subprocess.DEVNULL)

    writer = PdfWriter()
    for event, attendee, name in built:
        reader = PdfReader(os.path.join(PDF_DIR, name + ".pdf"))
        start = len(writer.pages)
        for page in reader.pages:
            writer.add_page(page)
        writer.add_outline_item(f"{event['title']} — {attendee['name']}", start)
    writer.add_metadata({"/Title": "Invitation letter requests — independently prepared drafts (6)",
                         "/Subject": "Not official organizer documents"})
    with open(os.path.join(ROOT, "Combined_All-Six-Invitation-Requests.pdf"), "wb") as f:
        writer.write(f)
    print("\n".join(n for _, _, n in built))


if __name__ == "__main__":
    sys.exit(main())
